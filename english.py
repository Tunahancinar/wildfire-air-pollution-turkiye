from pathlib import Path
import re
import io
import json
import zipfile

def tr(value):
    return value
import pandas as pd

def frame(data):
    if isinstance(data, pd.DataFrame):
        out = data.copy()
        for c in out:
            if pd.api.types.is_object_dtype(out[c]) or pd.api.types.is_string_dtype(out[c]):
                out[c] = out[c].map(tr)
        out.columns = [tr(c) for c in out.columns]
        out.index.names = [tr(n) for n in out.index.names]
        return out
    return tr(data)

def filename(name):
    return name

def workbook_bytes(path):
    from openpyxl import load_workbook
    wb = load_workbook(path)
    for ws in wb:
        if ws.title == 'Summary':
            ws.delete_rows(1, ws.max_row)
            ws.append(['Summary'])
            for line in summary_text(Path(path).parent).splitlines():
                ws.append([line])
        ws.title = tr(ws.title)[:31]
        for row in ws:
            for cell in row:
                if isinstance(cell.value, str):
                    cell.value = tr(cell.value)
        for col in ws.columns:
            ws.column_dimensions[col[0].column_letter].width = min(65, max(12, max((len(str(c.value or '')) for c in col[:100])) + 2))
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()

def english_zip(folder):
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, 'w', zipfile.ZIP_DEFLATED) as z:
        for p in sorted(Path(folder).rglob('*')):
            if not p.is_file():
                continue
            relative = filename(str(p.relative_to(Path(folder).parent)).replace('\\', '/'))
            if p.suffix == '.xlsx':
                data = workbook_bytes(p)
            elif p.suffix == '.csv':
                data = frame(pd.read_csv(p)).to_csv(index=False).encode('utf-8-sig')
            elif p.name == 'summary.txt':
                data = summary_text(p.parent).encode('utf-8')
            elif p.suffix in ('.txt', '.md'):
                data = tr(p.read_text(encoding='utf-8')).encode('utf-8')
            elif p.suffix == '.json':

                def convert(v):
                    if isinstance(v, dict):
                        return {tr(k): convert(x) for k, x in v.items()}
                    if isinstance(v, list):
                        return [convert(x) for x in v]
                    return tr(v)
                data = json.dumps(convert(json.loads(p.read_text(encoding='utf-8'))), ensure_ascii=False, indent=2).encode('utf-8')
            else:
                data = p.read_bytes()
            z.writestr(relative, data)
    return buf.getvalue()

def summary_text(folder):
    """Implementation of summary_text."""
    folder = Path(folder)
    summary = json.loads((folder / 'result.json').read_text(encoding='utf-8'))
    tables = pd.read_excel(folder / f'{folder.name}_report.xlsx', sheet_name=None)
    lines = [f'{folder.name} — WILDFIRE GAS ANALYSIS', f"Fire period: {summary['Start']} – {summary['End']} (source: {tr(summary.get('Date source', '—'))})", f"Combined gas signal: {summary.get('Combined gas signal', '—')}", '']
    for _, row in tables.get('Change detection', pd.DataFrame()).iterrows():
        gas = row['gas']
        onset = pd.to_datetime(row.get('increase_onset'), errors='coerce')
        if pd.isna(onset):
            lines.append(f'• {gas}: no event met the change threshold in the search period.')
            continue
        direction = 'increase' if row['direction'] == 'Increase' else 'decrease'
        date = pd.to_datetime(row.get('peak_date'), errors='coerce')
        date = date.strftime('%d.%m.%Y') if pd.notna(date) else 'not available'
        unit = summary.get(f'{gas} unit', 'μg/m3')
        lines.append(f"• {gas}: {direction} detected on {onset:%d.%m.%Y}; {int(row['agreeing_method'])} of 3 methods agree (confidence: {tr(row['confidence'])}).")
        lines.append(f"  {('Peak' if direction == 'increase' else 'Minimum')}: {row['peak_value']:.4g} {unit} on {date}; baseline: {row['baseline_median']:.4g} {unit}.")
        if pd.notna(row['increase_percent']):
            lines.append(f"  Change from baseline: {row['increase_percent']:+.1f}%.")
        if pd.notna(row.get('duration_day')):
            lines.append(f"  Duration of baseline departure: {int(row['duration_day'])} days.")
    perf = tables.get('Performance', pd.DataFrame())
    if not perf.empty:
        lines.extend(['', 'Performance (converted satellite vs. ground station):'])
        for _, row in perf.iterrows():
            if pd.notna(row.get('MAE')):
                lines.append(f"  {row['Gas']}: MAE={row['MAE']:.2f} μg/m3, RRMSE={row['RRMSE_%']:.2f}%, IA={row['IA']:.2f} (n={int(row['n'])}).")
    old = (folder / 'summary.txt').read_text(encoding='utf-8')
    if 'Notes:' in old:
        lines.extend(['', 'Notes:', tr(old.split('Notes:', 1)[1]).strip()])
    elif 'Notes:' in old:
        lines.extend(['', 'Notes:', old.split('Notes:', 1)[1].strip()])
    return '\n'.join(lines)

def install_plot_language():
    pass

def cli_text(text):
    return text

def install_cli_language():
    pass

class EnglishUI:
    """Implementation of EnglishUI."""
    TEXT = {'title', 'header', 'subheader', 'caption', 'text', 'markdown', 'write', 'info', 'warning', 'error', 'success', 'code', 'spinner', 'status', 'metric', 'expander', 'button', 'download_button', 'checkbox', 'toggle', 'slider', 'number_input', 'text_input', 'file_uploader', 'selectbox', 'radio', 'tabs', 'update', 'set_page_config'}

    def __init__(self, ui):
        self._ui = ui

    def __enter__(self):
        self._ui.__enter__()
        return self

    def __exit__(self, *args):
        return self._ui.__exit__(*args)

    def __getattr__(self, name):
        obj = getattr(self._ui, name)
        if name == 'sidebar':
            return EnglishUI(obj)
        if name not in self.TEXT | {'columns', 'empty', 'dataframe', 'table'}:
            return obj

        def call(*args, **kwargs):
            a = list(args)
            if name in {'selectbox', 'radio'}:
                original = kwargs.get('format_func', str)
                kwargs['format_func'] = lambda value: tr(original(value))
            if name in {'dataframe', 'table'}:
                if a:
                    a[0] = frame(a[0])
                if 'data' in kwargs:
                    kwargs['data'] = frame(kwargs['data'])
                if 'column_config' in kwargs:
                    kwargs['column_config'] = {tr(k): v for k, v in kwargs['column_config'].items()}
            elif name == 'tabs':
                a[0] = [tr(x) for x in a[0]]
            elif name in self.TEXT:
                limit = 3 if name in {'metric', 'write'} else 1
                for i in range(min(limit, len(a))):
                    a[i] = tr(a[i])
            for k in ['label', 'help', 'body', 'value' if name == 'metric' else '__unused', 'delta', 'page_title', 'placeholder']:
                if k in kwargs:
                    kwargs[k] = tr(kwargs[k])
            if name == 'download_button':
                if 'file_name' in kwargs:
                    kwargs['file_name'] = filename(kwargs['file_name'])
                if len(a) > 2:
                    a[2] = filename(a[2])
            result = obj(*a, **kwargs)
            if name in {'columns', 'tabs'}:
                return [EnglishUI(x) for x in result]
            if name in {'status', 'empty', 'expander'}:
                return EnglishUI(result)
            return result
        return call
