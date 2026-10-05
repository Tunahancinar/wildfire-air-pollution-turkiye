"""Implementation of the wildfire gas analysis module."""
from pathlib import Path
import json
from dataclasses import fields
import numpy as np
import pandas as pd
from scipy.stats import mannwhitneyu
from .config import AnalysisSettings, GASES
from .detection import detect_change_onset
from .analysis_statistics import make_periods, period_difference
NOTE = 'Confidence classes describe agreement between methods within one day, not statistical confidence probabilities. Mann–Whitney is a two-sided before-versus-during comparison. Small samples are flagged. P-values are unadjusted and exploratory; daily independence has not been verified. No detection and unavailable estimates are shown separately. Saved daily Excel series are preferred, with CSV fallback. Boundaries and conversion factors are not recalculated.'

def table_generate(root):
    rows = []
    allowed = {f.name for f in fields(AnalysisSettings)}
    for result in sorted(Path(root).glob('*/result.json')):
        folder = result.parent
        meta = json.loads(result.read_text(encoding='utf-8'))
        settings_path = folder / 'settings.json'
        settings = json.loads(settings_path.read_text(encoding='utf-8')) if settings_path.exists() else {}
        config = AnalysisSettings(**{k: v for k, v in settings.items() if k in allowed})
        start = pd.to_datetime(meta['Start'], dayfirst=True)
        end = pd.to_datetime(meta['End'], dayfirst=True)
        raw_mode = config.conversion == 'none'
        sheet = 'Series mol m-2' if raw_mode else 'Series ug m-3'
        csv = 'time_series_mol_m2.csv' if raw_mode else 'time_series_ug_m3.csv'
        report = folder / f'{folder.name}_report.xlsx'
        source = None
        error = None
        try:
            if report.exists() and sheet in pd.ExcelFile(report).sheet_names:
                df = pd.read_excel(report, sheet_name=sheet)
                source = 'Excel: ' + sheet
            elif (folder / csv).exists():
                df = pd.read_csv(folder / csv)
                source = 'CSV (saved precision)'
            else:
                df = pd.DataFrame()
                error = 'Saved daily observations unavailable'
            if not df.empty:
                df = df.set_index(df.columns[0])
                df.index = pd.to_datetime(df.index)
                df = df.sort_index().apply(pd.to_numeric, errors='coerce').replace([np.inf, -np.inf], np.nan)
                if not df.index.is_unique:
                    raise ValueError('Duplicate observations for the same day')
        except Exception as exc:
            df = pd.DataFrame()
            error = str(exc)
        periods = make_periods(start, end, config.before_day, config.after_day)
        fallback_gases = set()
        if not raw_mode:
            raw = pd.DataFrame()
            if report.exists() and 'Series mol m-2' in pd.ExcelFile(report).sheet_names:
                raw = pd.read_excel(report, sheet_name='Series mol m-2')
            elif (folder / 'time_series_mol_m2.csv').exists():
                raw = pd.read_csv(folder / 'time_series_mol_m2.csv')
            if not raw.empty:
                raw = raw.set_index(raw.columns[0])
                raw.index = pd.to_datetime(raw.index)
                raw = raw.sort_index().apply(pd.to_numeric, errors='coerce').replace([np.inf, -np.inf], np.nan)
                for gas in GASES:
                    if gas in raw and (gas not in df or df[gas].notna().sum() < 3):
                        df[gas] = raw[gas].reindex(df.index) if not df.empty else raw[gas]
                        fallback_gases.add(gas)
        differences = period_difference(df, periods).set_index('Gas') if not df.empty else pd.DataFrame()
        for gas, info in GASES.items():
            row = {'Fire': folder.name, 'Year': start.year, 'Fire start': start.strftime('%d.%m.%Y'), 'Fire end': end.strftime('%d.%m.%Y'), 'Gas': gas, 'Unit': 'mol/m2' if raw_mode else 'μg/m3', 'Data source': source, 'Threshold date': None, 'CUSUM onset': None, 'PELT onset': None, 'Consensus onset': None, 'Confidence class': 'Not estimable', 'Agreeing method count': 0, 'Direction': None, 'Delay (days)': None, 'Baseline observations': 0, 'Before n': 0, 'During n': 0, 'After n': 0, 'Mann–Whitney U': None, 'Mann–Whitney p': None, 'p<0.05 (unadjusted)': None, 'Test status': error or 'No converted observations for this gas'}
            if gas in df and df[gas].notna().sum() >= 3:
                if gas in fallback_gases:
                    row['Unit'] = 'mol/m2'
                    row['Data source'] = 'Saved satellite column series (no station conversion)'
                series = df[gas]
                detection, _ = detect_change_onset(series, gas, info.direction, start, end, config)
                for label, value in [('Threshold date', detection.threshold_date), ('CUSUM onset', detection.cusum_date), ('PELT onset', detection.pelt_date), ('Consensus onset', detection.increase_onset)]:
                    row[label] = value.strftime('%d.%m.%Y') if value is not None else 'Not detected'
                row.update({'Confidence class': detection.confidence, 'Agreeing method count': detection.agreeing_method, 'Direction': detection.direction, 'Delay (days)': detection.delay_day, 'Baseline observations': detection.baseline_n, 'Baseline source': detection.baseline_source})
                comparison = differences.loc[gas]
                for label in ['Before n', 'During n', 'After n', 'Before mean', 'During mean', 'After mean']:
                    row[label] = comparison[label]
                pre = series.loc[periods['Before'][0]:periods['Before'][1]].dropna()
                during = series.loc[start:end].dropna()
                if len(pre) >= 1 and len(during) >= 1:
                    test = mannwhitneyu(during, pre, alternative='two-sided')
                    row.update({'Mann–Whitney U': float(test.statistic), 'Mann–Whitney p': float(test.pvalue), 'p<0.05 (unadjusted)': 'Yes' if test.pvalue < 0.05 else 'No', 'Test status': 'Exploratory; small sample (<3 days)' if min(len(pre), len(during)) < 3 else 'Exploratory test'})
                    if np.isfinite(comparison['Mann-Whitney p']):
                        assert np.isclose(test.pvalue, comparison['Mann-Whitney p'])
                else:
                    row['Test status'] = 'Insufficient observations in the before or during period'
            elif gas in df:
                row['Test status'] = 'Fewer than 3 valid observations for detection'
            rows.append(row)
    table = pd.DataFrame(rows)
    if table.empty:
        return pd.DataFrame(columns=['Fire', 'Year', 'Gas', 'CUSUM onset', 'PELT onset', 'Consensus onset', 'Confidence class', 'Mann–Whitney U', 'Mann–Whitney p', 'Test status'])
    if not table.empty:
        first = ['Fire', 'Gas', 'CUSUM onset', 'PELT onset', 'Consensus onset', 'Confidence class', 'Mann–Whitney U', 'Mann–Whitney p', 'Test status']
        table = table[first + [c for c in table.columns if c not in first]]
    return table

def save(root):
    root = Path(root)
    root.mkdir(parents=True, exist_ok=True)
    table = table_generate(root)
    table.to_csv(root / 'ALL_FIRES_DETECTION_AND_MANN_WHITNEY.csv', index=False, encoding='utf-8-sig')
    with pd.ExcelWriter(root / 'ALL_FIRES_DETECTION_AND_MANN_WHITNEY.xlsx', engine='openpyxl') as writer:
        table.to_excel(writer, sheet_name='All fires', index=False)
        pd.DataFrame({'Method description': [NOTE]}).to_excel(writer, sheet_name='Method', index=False)
        ws = writer.sheets['All fires']
        ws.freeze_panes = 'F2'
        ws.auto_filter.ref = ws.dimensions
        from openpyxl.styles import Font, PatternFill
        for cell in ws[1]:
            cell.font = Font(bold=True, color='FFFFFF')
            cell.fill = PatternFill('solid', fgColor='18354B')
        for column in ws.columns:
            ws.column_dimensions[column[0].column_letter].width = min(48, max(15, max((len(str(c.value or '')) for c in column)) + 2))
        p_column = table.columns.get_loc('Mann–Whitney p') + 1
        for cells in ws.iter_rows(min_row=2):
            cells[p_column - 1].number_format = '0.000000'
    for name, group in table.groupby('Fire'):
        group.to_csv(root / name / 'detection_ve_mann_whitney.csv', index=False, encoding='utf-8-sig')
    return table
if __name__ == '__main__':
    import sys
    result = save(sys.argv[1] if len(sys.argv) > 1 else 'output')
    print(f"{result.Fire.nunique()} fire, {len(result)} gas rows; {result['Mann–Whitney p'].notna().sum()} computed tests")
